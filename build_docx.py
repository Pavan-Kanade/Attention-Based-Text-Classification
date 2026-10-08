import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def create_survey_paper():
    doc = Document()

    # Page Margins (IEEE Style: 0.75 inch)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # Styles
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Times New Roman'
    font.size = Pt(10)
    font.color.rgb = RGBColor(0, 0, 0)

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("Attention-Based Text Classification: A Survey of Attention Mechanisms for Deep Learning Based Text Classification")
    run_title.font.name = 'Times New Roman'
    run_title.font.size = Pt(20)
    run_title.font.bold = True

    doc.add_paragraph()

    # Authors
    p_authors = doc.add_paragraph()
    p_authors.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    r_a1 = p_authors.add_run("Abhijit Kamasane (34131), Vedant Kambale (34132), Pavan Kanade (34133), Student 4 (34134)\n")
    r_a1.font.bold = True
    r_a1.font.size = Pt(11)

    r_aff = p_authors.add_run("Department of Computer Science & Engineering, Third Year B.Tech\nProject-Based Learning (PBL) Deep Learning Submission\nStatus: Not Published (Pending Manual Confirmation)")
    r_aff.font.italic = True
    r_aff.font.size = Pt(9.5)

    doc.add_paragraph()

    # Abstract Box
    p_abs_title = doc.add_paragraph()
    r_abs_t = p_abs_title.add_run("Abstract—")
    r_abs_t.font.bold = True
    r_abs_t.font.italic = True
    r_abs_t.font.size = Pt(9)

    r_abs_body = p_abs_title.add_run(
        "Natural Language Processing (NLP) sequence classification models historically relied on Recurrent Neural Networks (RNNs) and Long Short-Term Memory (LSTM) networks. However, standard recurrent architectures compress variable-length input sequences into a single fixed-length hidden vector, creating a severe information bottleneck and offering minimal model interpretability. This survey paper examines the evolution, mathematical formulations, and practical applications of attention mechanisms for text classification. We analyze foundational additive attention (Bahdanau et al.), multiplicative dot-product attention (Luong et al.), multi-head self-attention (Vaswani et al.), and hierarchical document attention (Yang et al.). We present comparative analysis matrices, examine the architectural integration of Attention-BiLSTM networks, discuss real-world case studies in e-commerce feedback triage, and outline future research directions in interpretable deep learning."
    )
    r_abs_body.font.size = Pt(9)
    r_abs_body.font.italic = True

    p_kw = doc.add_paragraph()
    r_kw_t = p_kw.add_run("Index Terms—")
    r_kw_t.font.bold = True
    r_kw_t.font.italic = True
    r_kw_t.font.size = Pt(9)
    
    r_kw_body = p_kw.add_run("Attention Mechanism, Deep Learning, Text Classification, BiLSTM, Self-Attention, Query-Key-Value, Model Interpretability, IMDb Dataset.")
    r_kw_body.font.size = Pt(9)

    doc.add_paragraph()

    # Heading Helper
    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12)
        run.font.bold = True
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(10)
        run.font.bold = True
        run.font.italic = True
        return p

    def add_p(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(10)
        return p

    # I. INTRODUCTION
    add_h1("I. INTRODUCTION")
    add_p(
        "Text classification is a fundamental task in Natural Language Processing (NLP), with broad applications spanning sentiment analysis, spam detection, news categorization, legal document parsing, and customer support ticket routing. Traditional machine learning methods relied heavily on manual feature engineering, bag-of-words representations, or TF-IDF metrics paired with linear classifiers. While effective for short texts, these traditional approaches failed to capture word order, syntactic context, or complex semantic relationships."
    )
    add_p(
        "With the advent of deep learning, recurrent architectures such as Recurrent Neural Networks (RNNs) and Long Short-Term Memory (LSTM) networks became the standard paradigm for sequence modeling. By maintaining hidden state vectors across timesteps, LSTMs process text sequentially and capture long-range dependencies. However, conventional LSTM classifiers extract only the final hidden state vector h_T to represent the entire input text. For lengthy documents or multi-sentence reviews, this creates a critical information bottleneck, as early token information degrades over time."
    )
    add_p(
        "To resolve this bottleneck, attention mechanisms were introduced into deep learning. Attention allows the model to dynamically inspect all intermediate hidden states h_i and compute normalized scalar weights α_i that quantify the relative importance of each token. This paper surveys attention mechanisms for text classification, presenting mathematical formulations, literature comparisons, research gap analyses, and real-world case studies."
    )

    # II. BACKGROUND
    add_h1("II. BACKGROUND & SEQUENCE BOTTLENECK IN RECURRENT ARCHITECTURES")
    add_p(
        "In a standard unidirectional LSTM, an input sequence X = (x_1, x_2, ..., x_N) is mapped to a sequence of hidden state vectors (h_1, h_2, ..., h_N). In sequence classification, the standard approach feeds the final state h_N into a dense output layer. This design imposes two primary structural limitations:"
    )
    add_p(
        "1) Fixed-Length Bottleneck: The entire semantic content of a 500-word document must be compressed into a single vector (e.g., dimension 128), leading to information loss for early tokens."
    )
    add_p(
        "2) Lack of Interpretability: Standard LSTM hidden vectors function as opaque black boxes. Evaluators and users cannot determine which specific words caused a positive or negative sentiment classification."
    )

    # III. MATHEMATICAL FORMULATION
    add_h1("III. THE ATTENTION MECHANISM: MATHEMATICAL FORMULATION")
    add_p(
        "The additive attention mechanism proposed by Bahdanau et al. solves the bottleneck problem by computing alignment scores score_i between hidden states and a learnable context vector. For a sequence of BiLSTM hidden states h_i ∈ R^(2d_h), the attention layer computes:"
    )
    add_p("1) Attention Score Computation:\n   score_i = tanh(W_h · h_i + b)")
    add_p("2) Softmax Weight Normalization:\n   alpha_i = exp(score_i) / Σ_{j=1}^N exp(score_j)")
    add_p("3) Weighted Context Vector Construction:\n   c = Σ_{i=1}^N (alpha_i · h_i)")
    add_p(
        "The resulting context vector c is a weighted linear combination of all sequence hidden states, where tokens with higher scalar weights α_i contribute more strongly to the final classification head."
    )

    # IV. ATTENTION IN RECURRENT ARCHITECTURES
    add_h1("IV. ATTENTION IN RECURRENT ARCHITECTURES (LSTM & BiLSTM)")
    add_p(
        "Integrating attention with Bidirectional LSTM (BiLSTM) networks yields significant synergy. While unidirectional LSTMs only process past tokens (left-to-right), BiLSTMs compute both forward hidden states h_forward and backward hidden states h_backward at each timestep i, yielding concatenated representation h_i = [h_forward; h_backward]. The attention layer evaluates these bidirectional context states, ensuring that both preceding and succeeding context influence the token importance weight α_i."
    )

    # V. SCALED DOT-PRODUCT & Q-K-V
    add_h1("V. SCALED DOT-PRODUCT SELF-ATTENTION & QUERY-KEY-VALUE PARADIGM")
    add_p(
        "Vaswani et al. (2017) formalized attention using the Query-Key-Value (Q-K-V) abstraction. Given query matrix Q ∈ R^(N × d_k), key matrix K ∈ R^(N × d_k), and value matrix V ∈ R^(N × d_v), the Scaled Dot-Product Attention equation is formulated as:"
    )
    add_p("   Attention(Q, K, V) = softmax( (Q · Kᵀ) / √d_k ) · V")
    add_p(
        "Here, dot products of queries and keys measure alignment. The scaling factor 1/√d_k prevents large dot-product magnitudes from pushing the softmax function into regions with near-zero gradients."
    )

    # VI. TRANSFORMER VS SEQUENCE CLASSIFICATION
    add_h1("VI. TRANSFORMER-BASED ATTENTION VS SEQUENCE CLASSIFICATION ATTENTION")
    add_p(
        "It is essential to distinguish between sequence classification attention (used in Attention-BiLSTM models) and Transformer self-attention. Sequence classification attention uses a global attention parameter to aggregate BiLSTM hidden states into a single context vector c for dense classification. Transformer self-attention computes an N × N pairwise token relationship matrix, replacing recurrent timesteps entirely with multi-head attention blocks."
    )

    # VII. ATTENTION-BASED TEXT CLASSIFICATION SYSTEMS
    add_h1("VII. ATTENTION-BASED TEXT CLASSIFICATION SYSTEMS")
    add_p(
        "An end-to-end Attention-BiLSTM text classification system comprises: 1) Text Sanitization & Tokenization, 2) Word Embedding Layer, 3) Bidirectional LSTM Layer, 4) Additive Attention Layer, 5) Dense Classification Layer, and 6) Visual Token Attention Heatmap Renderer. In IMDb movie review sentiment benchmarks (50,000 samples), Attention-BiLSTM models achieve high test accuracy (~92.4%) while maintaining token-level visual interpretability."
    )

    # VIII. LITERATURE REVIEW
    add_h1("VIII. COMPREHENSIVE LITERATURE REVIEW")
    add_p("This section surveys four foundational research papers in attention mechanisms:")

    add_h2("A. Bahdanau et al. (2014) - Neural Machine Translation by Jointly Learning to Align and Translate")
    add_p("Pioneered additive attention in sequence-to-sequence translation, allowing decoders to dynamically align with encoder hidden states. Addresses the fixed-length vector bottleneck in long sequences.")

    add_h2("B. Luong et al. (2015) - Effective Approaches to Attention-based Neural Machine Translation")
    add_p("Introduced global and local attention mechanisms and compared dot-product, general, and location-based scoring functions, demonstrating matrix dot-product efficiency.")

    add_h2("C. Vaswani et al. (2017) - Attention Is All You Need")
    add_p("Introduced the Transformer architecture based entirely on multi-head scaled dot-product self-attention, eliminating recurrent LSTMs and enabling massive parallel training.")

    add_h2("D. Yang et al. (2016) - Hierarchical Attention Networks for Document Classification")
    add_p("Proposed hierarchical word-level and sentence-level attention networks for document classification on Yelp and IMDb review datasets, demonstrating that token heatmaps highlight core sentiment words.")

    # IX. COMPARATIVE TABLE
    add_h1("IX. COMPARATIVE LITERATURE MATRIX & DISCUSSION")
    
    table = doc.add_table(rows=5, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'

    headers = ["Paper / Reference", "Year", "Attention Mechanism", "Dataset / Task", "Main Contribution"]
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        p = hdr_cells[i].paragraphs[0]
        p.runs[0].font.bold = True
        p.runs[0].font.size = Pt(9)

    data = [
        ["Bahdanau et al.", "2014", "Additive Attention", "WMT 14 En-Fr", "Pioneered dynamic sequence alignment."],
        ["Luong et al.", "2015", "Dot-Product / Local", "WMT 14 NMT", "Simplified dot-product attention scoring."],
        ["Vaswani et al.", "2017", "Scaled Dot-Product (Q-K-V)", "WMT 14 En-De", "Introduced Transformer & self-attention."],
        ["Yang et al.", "2016", "Hierarchical Attention", "Yelp & IMDb", "Word & sentence level document attention."]
    ]

    for row_idx, row_data in enumerate(data, start=1):
        row_cells = table.rows[row_idx].cells
        for col_idx, cell_value in enumerate(row_data):
            row_cells[col_idx].text = cell_value
            p = row_cells[col_idx].paragraphs[0]
            p.runs[0].font.size = Pt(8.5)

    doc.add_paragraph()

    # X. RESEARCH GAP
    add_h1("X. RESEARCH GAP & BOTTLENECK ANALYSIS")
    add_p(
        "Despite significant advances, traditional sequence models still present a research gap between raw classification accuracy and interpretability. Conventional deep learning classifiers act as black boxes, making them difficult to audit in critical applications. Attention mechanisms bridge this gap by offering explicit token importance weights, allowing human evaluators to inspect model focus."
    )

    # XI. CASE STUDY
    add_h1("XI. REAL-WORLD CASE STUDIES & INDUSTRY APPLICATIONS")
    add_p(
        "Real-world application domains include e-commerce review analysis, support ticket triage, spam filtering, and social media brand monitoring. In support ticket triage, high attention weights on negative tokens (e.g., 'defective', 'broken', 'slow') trigger automated escalation to senior support agents."
    )

    # XII. CHALLENGES
    add_h1("XII. CHALLENGES, INTERPRETABILITY & LIMITATIONS")
    add_p(
        "Primary limitations include: 1) Computational Overhead: Additive tanh score calculations increase matrix operations, 2) Vocabulary Cutoff: Out-of-vocabulary words map to unknown tokens (<UNK>), and 3) Attention vs Causality: Attention weights indicate correlation rather than guaranteed causal explanation."
    )

    # XIII. FUTURE DIRECTIONS
    add_h1("XIII. FUTURE RESEARCH DIRECTIONS")
    add_p(
        "Future research includes fine-tuning pretrained Transformer models (BERT, RoBERTa), extending classification from binary to 5-star ratings, supporting multilingual tokenizers, and deploying cloud microservices."
    )

    # XIV. CONCLUSION
    add_h1("XIV. CONCLUSION")
    add_p(
        "This survey has examined attention mechanisms for text classification, contrasting additive sequence attention with scaled dot-product self-attention. By combining Bidirectional LSTMs with additive attention layers, deep learning models eliminate single-vector bottlenecks while providing visual token heatmaps."
    )

    # PUBLICATION NOTICE
    p_pub = doc.add_paragraph()
    p_pub.paragraph_format.space_before = Pt(12)
    r_pub = p_pub.add_run("PUBLICATION STATUS NOTICE: Not Published (Pending Manual Confirmation)")
    r_pub.font.bold = True
    r_pub.font.italic = True
    r_pub.font.color.rgb = RGBColor(180, 0, 0)

    # REFERENCES
    add_h1("REFERENCES")
    refs = [
        "[1] D. Bahdanau, K. Cho, and Y. Bengio, \"Neural machine translation by jointly learning to align and translate,\" arXiv preprint arXiv:1409.0473, 2014.",
        "[2] M.-T. Luong, H. Pham, and C. D. Manning, \"Effective approaches to attention-based neural machine translation,\" arXiv preprint arXiv:1508.04025, 2015.",
        "[3] A. Vaswani et al., \"Attention is all you need,\" in Advances in Neural Information Processing Systems (NeurIPS), 2017, pp. 5998–6008.",
        "[4] Z. Yang, D. Yang, C. Dyer, X. He, A. Smola, and E. Hovy, \"Hierarchical attention networks for document classification,\" in Proc. NAACL-HLT, 2016, pp. 1480–1489.",
        "[5] A. Maas, R. E. Daly, P. T. Pham, D. Huang, A. Y. Ng, and C. Potts, \"Learning word vectors for sentiment analysis,\" in Proc. ACL, 2011, pp. 142–150."
    ]
    for r in refs:
        add_p(r)

    out_path = "CIE_SUBMISSION/02_IEEE_Survey_Paper/Attention_Based_Text_Classification_IEEE_Survey_Paper.docx"
    doc.save(out_path)
    print(f"Generated {out_path}")

if __name__ == "__main__":
    create_survey_paper()
